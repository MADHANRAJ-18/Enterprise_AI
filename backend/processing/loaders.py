from __future__ import annotations

import io

import logging

from dataclasses import dataclass, field

from typing import List, Tuple

logger = logging.getLogger(__name__)

@dataclass

class PageInfo:
    page_number: int

    text: str

@dataclass

class LoadedDocument:
    full_text: str

    pages: List[PageInfo] = field(default_factory=list)

    page_count: int = 0

    file_type: str = ""

def _load_pdf(content: bytes) -> LoadedDocument:
    reader = None

    try:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(content), strict=False)

    except Exception:
        pass

    if reader is None:
        try:
            import PyPDF2

            reader = PyPDF2.PdfReader(io.BytesIO(content), strict=False)

        except Exception as exc:
            logger.warning("Failed to initialize PyPDF2 reader: %s", exc)

    pages: List[PageInfo] = []

    if reader is not None and hasattr(reader, "pages"):
        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""

            except Exception as exc:
                logger.warning("Page %d extraction failed: %s", i + 1, exc)

                text = ""

            pages.append(PageInfo(page_number=i + 1, text=text))

    full_text = "\n\n".join(p.text for p in pages if p.text.strip())

    if not full_text.strip():
        num_pages = len(pages) if pages else 1

        logger.info("PDF text stream empty — using page placeholders for %d page(s)", num_pages)

        pages = [PageInfo(page_number=i + 1, text=f"[Page {i + 1}: Image/Scanned content]") for i in range(num_pages)]

        full_text = "\n\n".join(p.text for p in pages)

    logger.info("PDF loaded: %d pages, %d chars", len(pages), len(full_text))

    return LoadedDocument(

        full_text=full_text,

        pages=pages,

        page_count=len(pages),

        file_type="PDF",

    )

def _load_docx(content: bytes) -> LoadedDocument:
    try:
        from docx import Document

    except ImportError:
        raise ImportError("python-docx is required. Run: pip install python-docx")

    doc = Document(io.BytesIO(content))

    paragraphs: List[str] = []

    for para in doc.paragraphs:
        txt = para.text.strip()

        if txt:
            paragraphs.append(txt)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(

                cell.text.strip() for cell in row.cells if cell.text.strip()

            )

            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)

    logger.info("DOCX loaded: %d paragraphs, %d chars", len(paragraphs), len(full_text))

    pages = [PageInfo(page_number=0, text=full_text)]

    return LoadedDocument(

        full_text=full_text,

        pages=pages,

        page_count=1,

        file_type="DOCX",

    )

def _load_txt(content: bytes) -> LoadedDocument:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            full_text = content.decode(encoding)

            break

        except UnicodeDecodeError:
            continue

    else:
        full_text = content.decode("utf-8", errors="replace")

    logger.info("TXT loaded: %d chars", len(full_text))

    pages = [PageInfo(page_number=0, text=full_text)]

    return LoadedDocument(

        full_text=full_text,

        pages=pages,

        page_count=1,

        file_type="TXT",

    )

def _load_pptx(content: bytes) -> LoadedDocument:
    try:
        from pptx import Presentation
        from pptx.util import Pt

    except ImportError:
        raise ImportError("python-pptx is required. Run: pip install python-pptx")

    prs = Presentation(io.BytesIO(content))

    pages: List[PageInfo] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_texts: List[str] = []

        for shape in slide.shapes:
            # Text frames (titles, body, text boxes)
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = para.text.strip()
                    if txt:
                        slide_texts.append(txt)

            # Tables embedded in slides
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        slide_texts.append(row_text)

        slide_text = "\n".join(slide_texts)
        pages.append(PageInfo(page_number=slide_num, text=slide_text))

    full_text = "\n\n".join(p.text for p in pages if p.text.strip())

    logger.info(
        "PPTX loaded: %d slides, %d chars",
        len(prs.slides),
        len(full_text),
    )

    return LoadedDocument(
        full_text=full_text,
        pages=pages,
        page_count=len(pages),
        file_type="PPTX",
    )


def load_document(content: bytes, file_type: str) -> LoadedDocument:
    ft = file_type.upper().strip()

    try:
        if ft == "PDF":
            return _load_pdf(content)

        elif ft == "DOCX":
            return _load_docx(content)

        elif ft == "TXT":
            return _load_txt(content)

        elif ft == "PPTX":
            return _load_pptx(content)

        else:
            raise ValueError(f"Unsupported file type: '{file_type}'. Expected PDF, DOCX, TXT, or PPTX.")

    except (ValueError, ImportError):
        raise

    except Exception as exc:
        logger.error("Text extraction failed for %s: %s", file_type, exc, exc_info=True)

        raise RuntimeError(f"Failed to extract text from {file_type}: {exc}") from exc
